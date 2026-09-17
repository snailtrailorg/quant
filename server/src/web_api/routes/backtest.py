"""回测 + 标的池 + Broker 用量 · 路由"""

from __future__ import annotations
import json
import asyncio
import logging
import os
import redis
from fastapi import APIRouter, Depends, Request, Body, WebSocket, WebSocketDisconnect
from ..auth import require_role, require_perm, audit_log
from ..errors import ApiError
from ..models import (LoginReq, UserCreate, StrategyConfig, InviteReq, RegisterReq, ForgotReq, ResetReq, ChangePwdReq, ChatReq, LLMModelReq, IMBotCreateReq, IMBotUpdateReq, IMBotUserReq, LlmBudgetReq, DataSourceReq, BrokerReq, RiskRuleReq, PoolReq, StrategyAccountReq)
from src.data_platform.db import get_conn, refresh_minute_symbols

logger = logging.getLogger("web_api")

_redis_pool = redis.ConnectionPool.from_url(
    os.environ.get("VALKEY_URL", "redis://127.0.0.1:6379/0"),
    decode_responses=True,
    socket_timeout=2, socket_connect_timeout=2)   # 批27-2：SSE gen() 内同步 get——挂起时帧断而非冻事件循环

_POOL_AGG_CACHE: dict[str, tuple[float, dict]] = {}   # 批27-26：minute-status 聚合 60s 缓存（键=pid）

router = APIRouter(tags=["backtest"])


@router.get("/api/pool")
def list_pools(payload: dict = Depends(require_perm("read"))):
    """标的池列表（含 symbols，#22）。"""
    with get_conn() as conn:
        try:
            conn.execute("SELECT 1 FROM pools LIMIT 1")
        except Exception:
            logger.warning("list_pools: pools 表不存在（需运行 alembic upgrade head）")
        try:
            conn.execute("SELECT 1 FROM pool_symbols LIMIT 1")
        except Exception:
            logger.warning("list_pools: pool_symbols 表不存在（需运行 alembic upgrade head）")
        cur = conn.execute(
            "SELECT p.id, p.name, p.category, p.description, ps.symbol, p.minute_history_start "
            "FROM pools p LEFT JOIN pool_symbols ps ON ps.pool_id=p.id ORDER BY p.id")
        rows = cur.fetchall()
        # minute_count：该池在展开表里的攒数据标的数（source='pool:{id}'，盲审 A-P2/B-P2）
        try:
            cur = conn.execute(
                "SELECT source, count(*) FROM minute_symbols WHERE source LIKE 'pool:%' GROUP BY source")
            pool_counts = {r[0][5:]: r[1] for r in cur.fetchall()}
        except Exception:
            pool_counts = {}
    pools = {}
    for pid, pname, pcat, pdesc, sym, mhs in rows:
        if pid not in pools:
            pools[pid] = {"id": pid, "name": pname, "category": pcat, "description": pdesc,
                          "symbols": [], "minute_history_start": str(mhs) if mhs else None,
                          "minute_count": pool_counts.get(str(pid), 0)}
        if sym:
            pools[pid]["symbols"].append(sym)
    return list(pools.values())


@router.post("/api/pool")
def create_pool(req: PoolReq, payload: dict = Depends(require_perm("strategy_control"))):
    """新建/更新标的池（#22）。"""
    symbols = [s.strip() for s in (req.symbolsStr or "").split("\n") if s.strip()]
    with get_conn() as conn:
        try:
            conn.execute("SELECT 1 FROM pools LIMIT 1")
        except Exception:
            logger.warning("create_pool: pools 表不存在（需运行 alembic upgrade head）")
        try:
            conn.execute("SELECT 1 FROM pool_symbols LIMIT 1")
        except Exception:
            logger.warning("create_pool: pool_symbols 表不存在（需运行 alembic upgrade head）")
        mhs = req.minute_history_start or None   # 空串/None → NULL（取消标记）
        conn.execute(
            "INSERT INTO pools (id, name, category, description, minute_history_start) VALUES (%s,%s,%s,%s,%s) "
            "ON CONFLICT (id) DO UPDATE SET name=EXCLUDED.name, category=EXCLUDED.category, "
            "description=EXCLUDED.description, "
            "minute_history_start=EXCLUDED.minute_history_start",
            (req.id, req.name, req.category, req.description, mhs))
        conn.execute("DELETE FROM pool_symbols WHERE pool_id=%s", (req.id,))
        for sym in symbols:
            conn.execute("INSERT INTO pool_symbols (pool_id, symbol) VALUES (%s,%s) ON CONFLICT DO NOTHING", (req.id, sym))
        conn.commit()
    refresh_minute_symbols()   # 池属性/成员变化 → 展开表同步（分钟数据源重构 21 号 §3.1）
    audit_log(payload["username"], "create_pool", req.id)
    return {"ok": True, "id": req.id, "count": len(symbols)}


@router.post("/api/pool/{pid}/symbol")
def add_pool_symbol_api(pid: str, body: dict = Body(...),
                        payload: dict = Depends(require_perm("strategy_control"))):
    """单标的入池（链条打磨：替代全量覆盖式 POST /api/pool——修并发覆盖竞态）。

    body: {symbol}——接受 vt 格式（600000.SHSE）或 Tushare 格式（600000.SH），归一到 vt。
    """
    from src.data_platform.schema import to_vt_symbol, vt_to_ts
    raw = (body.get("symbol") or "").strip()
    if not raw:
        raise ApiError(400, "MISSING_FIELDS", "symbol 必填")
    if "." not in raw:
        raise ApiError(400, "SYMBOL_INVALID", f"symbol 需带交易所后缀（如 600000.SHSE）: {raw}")
    vt = to_vt_symbol(raw)
    ts = vt_to_ts(vt)   # 校验可转换（防垃圾格式入池后同步空转——S-F1）
    with get_conn() as conn:
        cur = conn.execute("SELECT id, category FROM pools WHERE id=%s", (pid,))
        row = cur.fetchone()
        if not row:
            raise ApiError(404, "POOL_NOT_FOUND", f"池 {pid} 不存在")
        conn.execute(
            "INSERT INTO pool_symbols (pool_id, symbol) VALUES (%s, %s) "
            "ON CONFLICT (pool_id, symbol) DO NOTHING", (pid, vt))
        conn.commit()
    refresh_minute_symbols()   # 池成员变化 → 展开表同步
    # 二档深度数据回补（U 审项 9）：增量游标只认窗口，新标的的历史靠这一投——
    # 异步不阻塞响应；非 astock 池不投（pool_data 只拉 astock）
    if row[1] == "astock":
        from src.scheduler.tasks import pool_data_sync_task
        pool_data_sync_task.delay(symbols=[ts])
    audit_log(payload["username"], "pool_add_symbol", pid, vt)
    return {"status": "added", "symbol": vt, "ts_code": ts, "backfill": row[1] == "astock"}


@router.delete("/api/pool/{pid}/symbol/{sym}")
def del_pool_symbol_api(pid: str, sym: str,
                        payload: dict = Depends(require_perm("strategy_control"))):
    """单标的移出池。"""
    from src.data_platform.schema import to_vt_symbol
    vt = to_vt_symbol(sym)
    with get_conn() as conn:
        cur = conn.execute("DELETE FROM pool_symbols WHERE pool_id=%s AND symbol=%s RETURNING id",
                           (pid, vt))
        deleted = cur.fetchone()
        conn.commit()
    if not deleted:
        raise ApiError(404, "POOL_SYMBOL_NOT_FOUND", f"{vt} 不在池 {pid}")
    refresh_minute_symbols()   # 池成员变化 → 展开表同步
    audit_log(payload["username"], "pool_del_symbol", pid, vt)
    return {"status": "removed", "symbol": vt}


@router.get("/api/pool/{pid}/minute-status")
def pool_minute_status_api(pid: str,
                           payload: dict = Depends(require_perm("read"))):
    """池分钟数据覆盖状态（每标的 bar_1min 最后 ts——首轮回补可能 11.5h，进度可见是必须项）。

    批27-26：bar_1min 全表 GROUP BY 聚合按 pid 分键 60s 缓存（首轮回补期间高频轮询的 DB 压力点）。"""
    import time as _t
    _ck = f"minstat:{pid}"
    _hit = _POOL_AGG_CACHE.get(_ck)
    if _hit and _t.time() - _hit[0] < 60:
        return _hit[1]
    with get_conn() as conn:
        cur = conn.execute(
            "SELECT ps.symbol, COALESCE(b.last_ts::text, '') FROM pool_symbols ps "
            "LEFT JOIN (SELECT symbol, MAX(ts) AS last_ts FROM bar_1min GROUP BY symbol) b "
            "ON b.symbol = ps.symbol WHERE ps.pool_id=%s ORDER BY ps.symbol", (pid,))
        rows = cur.fetchall()
    _result = {"pool_id": pid, "symbols": [
        {"symbol": r[0], "last_ts": r[1][:19] if r[1] else None, "covered": bool(r[1])}
        for r in rows]}
    _POOL_AGG_CACHE[_ck] = (_t.time(), _result)
    return _result


@router.delete("/api/pool/{pid}")
def delete_pool(pid: str, payload: dict = Depends(require_perm("strategy_control"))):
    """删除标的池（CASCADE 删 symbols，#22）。"""
    with get_conn() as conn:
        conn.execute("DELETE FROM pools WHERE id=%s", (pid,))
        conn.commit()
    refresh_minute_symbols()   # 删整池 → 展开表移除该池成员（盲审 P1：漏了残留僵尸行）
    return {"ok": True}


@router.get("/api/minute-symbol")
def list_minute_symbols(payload: dict = Depends(require_perm("read"))):
    """攒数据标的展开表清单（分钟数据源重构 21 号 §3.5；含 last_ts 弱化漏取盲区）。"""
    with get_conn() as conn:
        cur = conn.execute(
            "SELECT m.symbol, m.source, COALESCE(b.last_ts::text, '') "
            "FROM minute_symbols m "
            "LEFT JOIN (SELECT symbol, MAX(ts) AS last_ts FROM bar_1min GROUP BY symbol) b "
            "ON b.symbol = m.symbol ORDER BY m.symbol")
        rows = cur.fetchall()
    return {"symbols": [
        {"symbol": r[0], "source": r[1], "last_ts": r[2][:19] if r[2] else None}
        for r in rows]}


@router.post("/api/minute-symbol/{symbol}")
def add_minute_symbol(symbol: str, payload: dict = Depends(require_perm("strategy_control"))):
    """个股直标攒分钟数据（source='direct'，UPSERT 覆盖池来源）。"""
    from src.data_platform.schema import to_vt_symbol
    raw = (symbol or "").strip()
    if not raw or "." not in raw:
        raise ApiError(400, "SYMBOL_INVALID", f"symbol 需带交易所后缀（如 600000.SHSE）: {raw}")
    vt = to_vt_symbol(raw)
    ex = vt.rsplit(".", 1)[-1].upper()
    if ex not in ("SHSE", "SZSE", "BSE"):
        raise ApiError(400, "SYMBOL_INVALID", f"不支持的交易所后缀: {ex}（支持 SHSE/SZSE/BSE）")
    with get_conn() as conn:
        conn.execute(
            "INSERT INTO minute_symbols (symbol, source) VALUES (%s, 'direct') "
            "ON CONFLICT (symbol) DO UPDATE SET source='direct', updated_at=now()", (vt,))
        conn.commit()
    audit_log(payload["username"], "minute_symbol_add", vt)
    return {"status": "added", "symbol": vt}


@router.delete("/api/minute-symbol/{symbol}")
def del_minute_symbol(symbol: str, payload: dict = Depends(require_perm("strategy_control"))):
    """取消个股直标（仅删 direct 行；删后 refresh 把仍属池的标的重物化为 pool 行）。"""
    from src.data_platform.schema import to_vt_symbol
    raw = (symbol or "").strip()
    if not raw or "." not in raw:
        raise ApiError(400, "SYMBOL_INVALID", f"symbol 需带交易所后缀（如 600000.SHSE）: {raw}")
    vt = to_vt_symbol(raw)
    with get_conn() as conn:
        conn.execute("DELETE FROM minute_symbols WHERE symbol=%s AND source='direct'", (vt,))
        conn.commit()
    refresh_minute_symbols()   # 盲审 A-P1/B-P1：add 的 ON CONFLICT 把 pool 行原地覆盖成 direct，
    # 删 direct 后标的虽仍在池、却从展开表消失——refresh 重物化 pool 行
    audit_log(payload["username"], "minute_symbol_del", vt)
    return {"status": "removed", "symbol": vt}


@router.delete("/api/backtest/{run_id}")
def delete_backtest(run_id: int, payload: dict = Depends(require_perm("strategy_control"))):
    """删除回测 run（backtest_symbols 随外键 ondelete=CASCADE 级联删）。

    造数脚本（seed-backtest/verify-gate）收尾清理测试 run，避免污染回测列表/成绩单。
    幂等：run 不存在返回 ok。盲审 P2：状态守卫（非终态 409 防 FK 违例）+ 审计。
    """
    with get_conn() as conn:
        cur = conn.execute("SELECT status FROM backtest_runs WHERE id=%s", (run_id,))
        row = cur.fetchone()
        if not row:
            conn.commit()
            return {"ok": True}   # 幂等：不存在返回 ok
        if row[0] in ("pending", "running"):
            conn.commit()
            raise ApiError(409, "RUN_ACTIVE", f"run {run_id} 仍在 {row[0]}，仅终态可删")
        conn.execute("DELETE FROM backtest_runs WHERE id=%s", (run_id,))
        conn.commit()
    audit_log(payload["username"], "backtest_delete", f"run {run_id}")
    return {"ok": True}


@router.post("/api/backtest")
def create_backtest_api(body: dict = Body(...),
                        payload: dict = Depends(require_perm("strategy_control"))):
    """启动回测 run：写 backtest_runs + Celery backtest_run_task。

    支持 symbol_params：per-symbol 参数覆盖。
    """
    strategy_id = body.get("strategy_config_id")
    symbols = body.get("symbols", [])
    pool_id = body.get("pool_id")
    if pool_id:
        with get_conn() as conn:
            cur = conn.execute("SELECT symbol FROM pool_symbols WHERE pool_id=%s", (pool_id,))
            symbols = [r[0] for r in cur.fetchall()]
    if not symbols or not strategy_id:
        raise ApiError(400, "MISSING_FIELDS", "需 strategy_config_id + symbols/pool_id")
    params = body.get("params", {})
    symbol_params = body.get("symbol_params", {})  # per-symbol 参数覆盖
    mode = body.get("mode", "single")
    # 批36a-3/4（盲审 A-P1-2 重写）：引擎资金键只读 run 级 params（tasks.py:927）——校验与执行
    # 同构分层：run 级 params 判保留键；symbol_params 每标的含保留键即 400（执行侧不消费=脏配置）
    # +defs 按执行同构 defaults⊕params⊕per_symbol 逐标的判定。
    from math import isfinite
    from src.strategy_framework.strategy import validate_params_against_defs
    if not isinstance(params, dict) or not isinstance(symbol_params, dict):
        raise ApiError(400, "BAD_BACKTEST_PARAM", "params/symbol_params 须为对象")
    def _chk_money(d: dict, where: str):
        try:
            capital = float(d.get("capital", 1_000_000))
            commission = float(d.get("commission", 0.0005))
            slippage = float(d.get("slippage", 0))
        except (TypeError, ValueError):
            raise ApiError(400, "BAD_BACKTEST_PARAM", f"{where} capital/commission/slippage 必须是数字")
        if not all(isfinite(x) for x in (capital, commission, slippage)):
            raise ApiError(400, "BAD_BACKTEST_PARAM", f"{where} 资金参数须为有限数字")
        if not (1e4 <= capital <= 1e10):
            raise ApiError(400, "BAD_BACKTEST_PARAM", "初始资金须在 1 万 ~ 100 亿之间")
        if not (0 <= commission <= 0.01):
            raise ApiError(400, "BAD_BACKTEST_PARAM", "佣金率须在 0 ~ 0.01（比例）之间（负佣金=回测造假面）")
        if not (0 <= slippage <= 0.05):
            raise ApiError(400, "BAD_BACKTEST_PARAM", "滑点须在 0 ~ 0.05 之间")
    _chk_money(params, "params")
    for sym, sp in symbol_params.items():
        if not isinstance(sp, dict):
            raise ApiError(400, "BAD_BACKTEST_PARAM", f"symbol_params[{sym}] 须为对象")
        if {"capital", "commission", "slippage"} & set(sp):
            raise ApiError(400, "BAD_BACKTEST_PARAM",
                           f"symbol_params[{sym}] 不可包含 capital/commission/slippage（引擎只读全局参数）")
    # 批36a-4：validate_params_against_defs 接线（现状唯一未接的写端点——strategy/live-task 已接）
    with get_conn() as conn:
        cur = conn.execute("SELECT params FROM strategy_config WHERE id=%s", (strategy_id,))
        srow = cur.fetchone()
        if not srow:
            raise ApiError(404, "STRATEGY_NOT_FOUND", f"策略 {strategy_id} 不存在")
    try:
        import json as _json
        sparams = srow[0] if isinstance(srow[0], dict) else _json.loads(srow[0] or "{}")
    except (ValueError, TypeError):
        sparams = {}
    defs = sparams.get("parameter_defs") or []
    merged = dict(params)   # defs 校验面：run 级（symbol_params 仅策略参数覆盖——逐标的并入判定）
    for sym in symbols:
        err = validate_params_against_defs({**merged, **(symbol_params.get(sym) or {})}, defs)
        if err:
            raise ApiError(400, "BAD_BACKTEST_PARAM", f"[{sym}] {err}")
    with get_conn() as conn:
        cur = conn.execute(
            "INSERT INTO backtest_runs (strategy_config_id, symbols, params, symbol_params, mode, status) "
            "VALUES (%s,%s,%s,%s,%s,'pending') RETURNING id",
            (strategy_id, json.dumps(symbols), json.dumps(params),
             json.dumps(symbol_params), mode))
        run_id = cur.fetchone()[0]
        conn.commit()
    from src.scheduler.tasks import backtest_run_task
    task = backtest_run_task.delay(run_id)
    audit_log(payload["username"], "backtest_create", f"run {run_id}")
    return {"run_id": run_id, "task_id": task.id}


@router.get("/api/broker-usage")
def broker_usage(payload: dict = Depends(require_perm("read"))):
    """通道调用量监控（#37，broker_usage 表聚合）。"""
    with get_conn() as conn:
        try:
            conn.execute("SELECT 1 FROM broker_usage LIMIT 1")
        except Exception:
            logger.warning("broker_usage: broker_usage 表不存在（需运行 alembic upgrade head）")
        cur = conn.execute(
            "SELECT provider, COUNT(*), COALESCE(AVG(latency_ms),0), "
            "CASE WHEN COUNT(*)>0 THEN round(SUM(CASE WHEN success THEN 1 ELSE 0 END)*100.0/COUNT(*),1) ELSE 0 END "
            "FROM broker_usage WHERE ts::date=current_date GROUP BY provider ORDER BY COUNT(*) DESC")
        today = [{"provider": r[0], "calls": r[1], "avg_latency_ms": int(r[2]), "success_rate": float(r[3])} for r in cur.fetchall()]
        cur = conn.execute(
            "SELECT ts::date AS d, COUNT(*), COALESCE(AVG(latency_ms),0) FROM broker_usage "
            "WHERE ts >= current_date - interval '7 days' GROUP BY d ORDER BY d")
        trend = [{"date": str(r[0]), "calls": r[1], "avg_latency_ms": int(r[2])} for r in cur.fetchall()]
    return {"today": today, "trend": trend}


def _safe_json(v, fallback):
    """H11（01 P0#4）：中断残行/引擎写坏的 JSON 不再炸整页 500——单行降级+留痕。"""
    if not v:
        return fallback
    try:
        return json.loads(v)
    except Exception:
        logger.warning("backtest JSON 字段解析失败（降级 %r）: %r", type(fallback).__name__, str(v)[:120])
        return fallback


@router.get("/api/backtest")
def list_backtest_api(payload: dict = Depends(require_perm("read"))):
    with get_conn() as conn:
        cur = conn.execute(
            "SELECT b.id, b.strategy_config_id, b.symbols, b.mode, b.status, b.created_at, b.finished_at, b.summary_metrics, "
            "b.task_id "
            "FROM backtest_runs b ORDER BY b.id DESC LIMIT 100")
        rows = cur.fetchall()
    return [{"id": r[0], "strategy_config_id": r[1], "symbols": _safe_json(r[2], []),
             "task_id": r[8],   # H11 根修(backtest_symbols 无 task_id 列——原子查询 UndefinedColumn=500 真因,改 runs 级)
             "mode": r[3], "status": r[4], "created_at": str(r[5]) if r[5] else None,
             "finished_at": str(r[6]) if r[6] else None,
             "summary": _safe_json(r[7], {})} for r in rows]


@router.get("/api/backtest/{run_id}")
def get_backtest_api(run_id: int,
                     payload: dict = Depends(require_perm("read"))):
    with get_conn() as conn:
        cur = conn.execute(
            "SELECT id, strategy_config_id, symbols, params, mode, status, summary_metrics, created_at, finished_at, task_id "
            "FROM backtest_runs WHERE id=%s", (run_id,))
        r = cur.fetchone()
        if not r:
            raise ApiError(404, "BACKTEST_NOT_FOUND", "run 不存在")
        # H11 同款修复（详情端点漏修——backtest_symbols 无 task_id 列，原子查询 UndefinedColumn=500）：
        # task_id 取 runs 级
        cur = conn.execute(
            "SELECT symbol, status, result FROM backtest_symbols WHERE run_id=%s ORDER BY symbol",
            (run_id,))
        syms = cur.fetchall()
    # 链条打磨#16（2026-08-19）：补前端实际读取的形状——顶层绩效四卡（此前恒 '-'）+
    # symbols 改对象数组（此前字符串数组致状态列空白）+ 顶层 task_id（终止按钮 #17）
    _mk = [_s for _s in syms if _s[1] == "done"]
    _agg = {}
    for k in ("total_return_pct", "win_rate", "max_drawdown_pct", "sharpe_ratio", "total_trades",
              "volatility", "sortino_ratio", "alpha", "beta", "information_ratio",
              "benchmark_return", "benchmark_volatility"):
        vals = [float((_safe_json(_s[2], {})).get(k) or 0) for _s in _mk if _s[2]]   # 批27-9：_safe_json 防残行炸详情页
        _agg[k] = round(sum(vals) / len(vals), 3) if vals else None
    _run_task = r[9]   # H11：runs 级 task_id
    # wd-20 §1.2 验证门派生字段（单点）：params 区间优先，回落 created_at→finished_at
    import datetime as _dt

    def _d(x):
        try:
            return _dt.date.fromisoformat(str(x)[:10])
        except Exception:
            return None

    _p = _safe_json(r[3], {})   # 批27-9
    _s0, _e0 = _d(_p.get("start")), _d(_p.get("end"))
    if _s0 and _e0:
        span_days = (_e0 - _s0).days + 1
    else:
        # 盲审B-P3：回退改 done 符号 result.start/end_date 聚合（引擎每符号都写）——
        # 度量回测窗口而非执行时长（分钟级 run 用 created→finished 会恒 1 天被 90 门误拦）
        _ss = [_d((_safe_json(_x[2], {})).get("start_date")) for _x in syms if _x[1] == "done" and _x[2]]   # 批27-9
        _ee = [_d((_safe_json(_x[2], {})).get("end_date")) for _x in syms if _x[1] == "done" and _x[2]]
        _ss = [x for x in _ss if x]; _ee = [x for x in _ee if x]
        if _ss and _ee:
            span_days = (max(_ee) - min(_ss)).days + 1
        else:
            span_days = 0
    # 批16 bug3：补 annualized_return（前端卡绑定此字段但引擎/聚合从无——恒 '—'）。
    # 口径 ((1+r/100)^(365/span)-1)*100；span<30 天复利外推会爆炸，钳制为 None（不做年化）。
    _annualized = None
    if span_days and span_days >= 30 and _agg.get("total_return_pct") is not None:
        try:
            _r = float(_agg["total_return_pct"]) / 100.0
            _annualized = round(((1 + _r) ** (365.0 / span_days) - 1) * 100, 1)
        except Exception:
            _annualized = None
    return {"id": r[0], "strategy_config_id": r[1],
            "span_days": span_days,
            "annualized_return": _annualized,
            "symbols": [{"symbol": _s[0], "status": _s[1],
                         "result": _safe_json(_s[2], {}) if _s[2] else {}} for _s in syms],   # 批27-9
            "symbols_list": _safe_json(r[2], []),
            "params": _safe_json(r[3], {}), "mode": r[4], "status": r[5],
            "summary": _safe_json(r[6], {}) if r[6] else {},
            "task_id": _run_task,
            "total_return_pct": _agg.get("total_return_pct"),
            "win_rate": _agg.get("win_rate"),
            "max_drawdown_pct": _agg.get("max_drawdown_pct"),
            "sharpe_ratio": _agg.get("sharpe_ratio"),
            "volatility": _agg.get("volatility"),
            "sortino_ratio": _agg.get("sortino_ratio"),
            "alpha": _agg.get("alpha"),
            "beta": _agg.get("beta"),
            "information_ratio": _agg.get("information_ratio"),
            "benchmark_return": _agg.get("benchmark_return"),
            "benchmark_volatility": _agg.get("benchmark_volatility"),
            "total_trades": _agg.get("total_trades"),
            "symbols_detail": [{"symbol": _s[0], "status": _s[1], "result": _safe_json(_s[2], {}) if _s[2] else {}}
                              for _s in syms]}


@router.get("/api/backtest/{run_id}/summary")
def backtest_summary(run_id: int, payload: dict = Depends(require_perm("read"))):
    """回测组汇总：标的绩效平均+排名（#22）。"""
    with get_conn() as conn:
        cur = conn.execute("SELECT symbol, result FROM backtest_symbols WHERE run_id=%s AND status='done'", (run_id,))
        rows = cur.fetchall()
    metrics_keys = ["total_return_pct", "win_rate", "max_drawdown_pct", "sharpe_ratio", "total_trades",
                    "volatility", "sortino_ratio", "alpha", "beta", "information_ratio",
                    "benchmark_return", "benchmark_volatility"]
    results = []
    for sym, result_json in rows:
        r = _safe_json(result_json, {})
        results.append({"symbol": sym, **{k: r.get(k, 0) for k in metrics_keys}})
    ranked = sorted(results, key=lambda x: x.get("total_return_pct", 0), reverse=True)
    avg = {k: round(sum(r[k] for r in results) / len(results), 3) for k in metrics_keys} if results else {}
    return {"run_id": run_id, "count": len(results), "avg": avg, "ranked": ranked}


# 滚动绩效 type → result.metrics.rolling 的指标键（ptrade 批 2）
_ROLLING_METRIC_KEYS = {
    "return": "return", "benchmark": "benchmark_return", "alpha": "alpha",
    "beta": "beta", "sharpe": "sharpe", "sortino": "sortino",
    "information": "information_ratio", "volatility": "volatility",
    "drawdown": "max_drawdown",
}


@router.get("/api/backtest/{run_id}/metrics")
def backtest_metrics(run_id: int, type: str = "return",
                     payload: dict = Depends(require_perm("read"))):
    """滚动绩效（ptrade 批 2）：按 type 返回月度 × 窗口（1/3/6/12）的指标二维表。

    type ∈ return|benchmark|alpha|beta|sharpe|sortino|information|volatility|drawdown。
    多标的按均值聚合。返回 {"data": {month: {window: value}}}，前端滚动二维表直接渲染。
    """
    key = _ROLLING_METRIC_KEYS.get(type)
    if key is None:
        raise ApiError(400, "INVALID_METRIC_TYPE",
                       f"type {type} 非法，须 ∈ {list(_ROLLING_METRIC_KEYS)}")
    with get_conn() as conn:
        cur = conn.execute("SELECT symbol, result FROM backtest_symbols WHERE run_id=%s AND status='done'", (run_id,))
        rows = cur.fetchall()
    agg: dict = {}
    for _sym, result_json in rows:
        r = _safe_json(result_json, {})
        rolling = (r.get("metrics") or {}).get("rolling") or {}
        for month, windows in rolling.items():
            agg.setdefault(month, {})
            for win, m in windows.items():
                if m is None:
                    continue
                v = m.get(key)
                if v is None:
                    continue
                agg[month].setdefault(win, []).append(v)
    data = {}
    for month in sorted(agg):
        data[month] = {}
        for win in ("1", "3", "6", "12"):
            vals = agg.get(month, {}).get(win)
            data[month][win] = round(sum(vals) / len(vals), 4) if vals else None
    return {"run_id": run_id, "type": type, "windows": ["1", "3", "6", "12"], "data": data}


@router.get("/api/backtest/{run_id}/export")
def backtest_export(run_id: int, symbol: str | None = None,
                    payload: dict = Depends(require_perm("read"))):
    """导出回测报告为 Excel（ptrade 批 3）：指标概览 / 交易明细 / 每日持仓 三 sheet。

    同步生成（数据量小，10s 内完成），返回 xlsx 文件流。symbol 可选（单标的导出，前端单标视图传）。
    """
    import io
    from openpyxl import Workbook
    from fastapi.responses import StreamingResponse

    sql = "SELECT symbol, result FROM backtest_symbols WHERE run_id=%s AND status='done'"
    params = [run_id]
    if symbol:
        sql += " AND symbol=%s"
        params.append(symbol)
    sql += " ORDER BY symbol"
    with get_conn() as conn:
        cur = conn.execute(sql, params)
        rows = cur.fetchall()
    if not rows:
        raise ApiError(404, "BACKTEST_NOT_FOUND", "run 无 done 结果")

    wb = Workbook()
    ws_meta = wb.active
    ws_meta.title = "指标概览"
    metric_cols = ["symbol", "total_return_pct", "benchmark_return", "alpha", "beta", "sharpe_ratio",
                   "sortino_ratio", "information_ratio", "volatility", "benchmark_volatility",
                   "max_drawdown_pct", "win_rate", "total_trades"]
    ws_meta.append(metric_cols)
    for sym, result_json in rows:
        r = _safe_json(result_json, {})
        ws_meta.append([sym] + [r.get(k, "") for k in metric_cols[1:]])

    ws_trades = wb.create_sheet("交易明细")
    ws_trades.append(["symbol", "ts", "action", "volume", "price", "commission"])
    for sym, result_json in rows:
        r = _safe_json(result_json, {})
        for t in (r.get("trades") or []):
            ws_trades.append([sym, t.get("ts"), t.get("action"), t.get("volume"), t.get("price"), t.get("commission")])

    ws_pos = wb.create_sheet("每日持仓")
    ws_pos.append(["symbol", "ts", "close", "position", "avg_price", "cash", "value"])
    for sym, result_json in rows:
        r = _safe_json(result_json, {})
        for d in (r.get("daily_values") or []):
            ws_pos.append([sym, d.get("ts"), d.get("close"), d.get("position"), d.get("avg_price"), d.get("cash"), d.get("value")])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=backtest_{run_id}.xlsx"})


# 回测报告 PDF 多语言列头（ptrade 批 3，N 语言：新增语言只加 key）
_PDF_COLUMNS = {
    "zh": {
        "meta": ["标的", "总收益%", "基准收益%", "阿尔法", "贝塔", "夏普", "索提诺", "信息率", "策略波动率%", "基准波动率%", "最大回撤%", "胜率%", "交易次数"],
        "trades": ["标的", "时间", "方向", "数量", "价格", "佣金"],
        "positions": ["标的", "日期", "收盘价", "持仓", "成本", "现金", "总资产"],
    },
    "en": {
        "meta": ["Symbol", "Total Return%", "Benchmark%", "Alpha", "Beta", "Sharpe", "Sortino", "Info Ratio", "Volatility%", "Bench Volatility%", "Max Drawdown%", "Win Rate%", "Trades"],
        "trades": ["Symbol", "Time", "Action", "Volume", "Price", "Commission"],
        "positions": ["Symbol", "Date", "Close", "Position", "Avg Cost", "Cash", "Total Value"],
    },
}


@router.get("/api/backtest/{run_id}/export.pdf")
def backtest_export_pdf(run_id: int, symbol: str | None = None, lang: str = "en",
                        payload: dict = Depends(require_perm("read"))):
    """导出回测报告为 PDF（ptrade 批 3）：指标概览/交易明细/每日持仓三表。

    weasyprint HTML→PDF：系统字体（wqy-microhei 等）自动 fc 发现，无需手动注册；
    文本按 lang（normalize_lang 回落 en）。
    """
    import html as _html
    import io
    from fastapi.responses import StreamingResponse
    from src.email_service import normalize_lang
    try:
        from weasyprint import HTML
    except Exception:
        raise ApiError(503, "PDF_EXPORT_UNAVAILABLE", "PDF 导出依赖 weasyprint 未装（服务器需 pango/cairo 系统库）")

    lang = normalize_lang(lang)
    # .get 回落 en：加语言只在 terms.py 登记而忘加 _PDF_COLUMNS key 时，不 KeyError→500，与 email 侧一致
    cols = _PDF_COLUMNS.get(lang, _PDF_COLUMNS["en"])

    sql = "SELECT symbol, result FROM backtest_symbols WHERE run_id=%s AND status='done'"
    params = [run_id]
    if symbol:
        sql += " AND symbol=%s"
        params.append(symbol)
    sql += " ORDER BY symbol"
    with get_conn() as conn:
        cur = conn.execute(sql, params)
        rows = cur.fetchall()
    if not rows:
        raise ApiError(404, "BACKTEST_NOT_FOUND", "run 无 done 结果")

    def _esc(v):
        return _html.escape(str(v) if v is not None else "")

    def _table(headers, data_rows):
        thead = "".join(f"<th>{_esc(h)}</th>" for h in headers)
        tbody = "".join("<tr>" + "".join(f"<td>{_esc(c)}</td>" for c in row) + "</tr>" for row in data_rows)
        return f"<table><thead><tr>{thead}</tr></thead><tbody>{tbody}</tbody></table>"

    meta_rows = []
    for sym, result_json in rows:
        r = _safe_json(result_json, {})
        meta_rows.append([sym, r.get("total_return_pct", ""), r.get("benchmark_return", ""),
                          r.get("alpha", ""), r.get("beta", ""), r.get("sharpe_ratio", ""),
                          r.get("sortino_ratio", ""), r.get("information_ratio", ""),
                          r.get("volatility", ""), r.get("benchmark_volatility", ""),
                          r.get("max_drawdown_pct", ""), r.get("win_rate", ""), r.get("total_trades", "")])

    trade_rows = []
    for sym, result_json in rows:
        r = _safe_json(result_json, {})
        for t in (r.get("trades") or []):
            trade_rows.append([sym, (t.get("ts") or "")[:19], t.get("action"), t.get("volume"), t.get("price"), t.get("commission")])

    pos_rows = []
    for sym, result_json in rows:
        r = _safe_json(result_json, {})
        for d in (r.get("daily_values") or []):
            pos_rows.append([sym, (d.get("ts") or "")[:10], d.get("close"), d.get("position"), d.get("avg_price"), d.get("cash"), d.get("value")])

    doc_html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8">
<style>
  body {{ font-family: "WenQuanYi Micro Hei", "Noto Sans CJK SC", sans-serif; font-size: 10px; }}
  table {{ border-collapse: collapse; width: 100%; margin-bottom: 18px; }}
  th, td {{ border: 1px solid #999; padding: 4px 6px; text-align: left; }}
  th {{ background: #eee; font-weight: bold; }}
</style></head>
<body>
{_table(cols["meta"], meta_rows)}
{_table(cols["trades"], trade_rows)}
{_table(cols["positions"], pos_rows)}
</body></html>"""

    buf = io.BytesIO()
    try:
        HTML(string=doc_html).write_pdf(buf)
    except Exception:
        raise ApiError(503, "PDF_EXPORT_FAIL", "PDF 渲染失败（字体/系统库缺失）")
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename=backtest_{run_id}.pdf"})


@router.get("/api/backtest/{run_id}/{symbol}/stream")
def backtest_stream_api(run_id: int, symbol: str,
                        payload: dict = Depends(require_perm("read"))):
    """SSE 单标的实时（轮询 Valkey backtest:run:{run_id}:{symbol}）。"""
    from fastapi.responses import StreamingResponse
    r = redis.Redis(connection_pool=_redis_pool)
    key = f"backtest:run:{run_id}:{symbol}"

    async def gen():
        for _ in range(720):  # 最多 6 分钟
            done = r.get(key + ":done")
            if done:
                yield f"data: {done}\n\n"
                break
            err = r.get(key + ":error")
            if err:
                yield f"data: {json.dumps({'error': err})}\n\n"
                break
            frame = r.get(key)
            if frame:
                yield f"data: {frame}\n\n"
            await asyncio.sleep(0.5)

    return StreamingResponse(gen(), media_type="text/event-stream")